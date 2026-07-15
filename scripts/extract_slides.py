#!/usr/bin/env python3
"""Extract a .pptx deck into markdown: per-slide text, speaker notes, and images.

Path 2 of the waterfall. Speaker notes are first-class here — for a tutorial they
often carry the explanation that the slide only gestures at, so they're the best
comprehension signal in the whole deck. Each slide is also rendered to an image so
diagrams and hand-drawn maths survive for the Q&A phase.

    python extract_slides.py deck.pptx --out materials/<slug>
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

# Common install locations for the LibreOffice headless binary.
SOFFICE_CANDIDATES = [
    "soffice", "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice", "/usr/local/bin/soffice",
]


def _find_soffice() -> str | None:
    for cand in SOFFICE_CANDIDATES:
        found = shutil.which(cand) if "/" not in cand else (cand if Path(cand).exists() else None)
        if found:
            return found
    return None


def _slide_title(slide) -> str:
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        return slide.shapes.title.text.strip()
    return ""


def _slide_body(slide) -> str:
    """All non-title text on the slide, in reading order, tables included."""
    title = slide.shapes.title
    parts: list[str] = []
    for shape in slide.shapes:
        if shape is title:
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append("| " + " | ".join(cells) + " |")
    return "\n\n".join(parts)


def _speaker_notes(slide) -> str:
    if slide.has_notes_slide:
        return slide.notes_slide.notes_text_frame.text.strip()
    return ""


def _render_images(pptx: Path, images_dir: Path, dpi: int = 150) -> int:
    """pptx -> pdf (LibreOffice headless) -> per-slide PNG (PyMuPDF). Returns count."""
    soffice = _find_soffice()
    if soffice is None:
        print("WARNING: LibreOffice not found; skipping slide images")
        return 0

    images_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        proc = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(pptx)],
            capture_output=True, text=True,
        )
        if proc.returncode != 0:
            print(f"WARNING: soffice failed, skipping images:\n{proc.stderr}")
            return 0
        pdfs = list(Path(tmp).glob("*.pdf"))
        if not pdfs:
            return 0

        import fitz  # PyMuPDF

        doc = fitz.open(pdfs[0])
        zoom = dpi / 72.0
        mat = fitz.Matrix(zoom, zoom)
        for i, page in enumerate(doc, start=1):
            page.get_pixmap(matrix=mat).save(images_dir / f"slide-{i}.png")
        return doc.page_count


def extract(pptx: Path, out: Path) -> dict:
    """Return {markdown_path, images_dir, n_slides, n_images}."""
    out.mkdir(parents=True, exist_ok=True)
    images_dir = out / "source" / "slides"

    prs = Presentation(str(pptx))
    n_images = _render_images(pptx, images_dir)

    lines: list[str] = [f"# {pptx.stem}\n"]
    n = 0
    for i, slide in enumerate(prs.slides, start=1):
        n = i
        title = _slide_title(slide)
        header = f"## Slide {i}" + (f": {title}" if title else "")
        lines.append(header)

        body = _slide_body(slide)
        if body:
            lines.append(body)

        notes = _speaker_notes(slide)
        if notes:
            lines.append(f"**Notes:** {notes}")

        img = images_dir / f"slide-{i}.png"
        if img.exists():
            rel = img.relative_to(out)
            lines.append(f"![slide {i}]({rel})")

        lines.append("")  # blank line between slides

    md_path = out / "material.md"
    md_path.write_text("\n".join(lines))
    return {"markdown_path": str(md_path), "images_dir": str(images_dir),
            "n_slides": n, "n_images": n_images}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    print(json.dumps(extract(args.pptx, args.out), indent=2))


if __name__ == "__main__":
    main()
